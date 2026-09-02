#!/usr/bin/env python3
"""Prepare auditable publication data from the locked v4 screening run."""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from typing import Any, Iterable, Mapping, Sequence
import uuid

import numpy as np


HERE = Path(__file__).resolve().parent
DEFAULT_REPO_ROOT = HERE.parents[1]
DEFAULT_OUTPUT_DIR = HERE / "data"
CONFIG_PATH = HERE / "config.json"
NUMERIC_SUFFIXES = ("_score", "_rank", "_percentile", "_cosine")
PUBLIC_OUTPUT_RECORDS = {
    "anchor_cosines.csv",
    "diagnostics.json",
    "semantic_scores.csv",
    "semantic_scores.jsonl",
    "top50.csv",
    "semantic_embeddings.npy",
    "embedding_completion.npy",
    "checkpoint_metadata.json",
}
WORKBOOK_SHEETS = {
    "Ranking_200": "annotated/ranking_all_annotated.csv",
    "Axis_scores_600": "anonymous/axis_scores_long.csv",
    "Top20_matrix": "anonymous/top20_axis_matrix.csv",
    "Top50": "anonymous/top50_ranking.csv",
    "Selected_anchor_30": "annotated/selected_anchor_profile.csv",
    "Stage1_800": "anonymous/stage1_all.csv",
    "Anchor_frequency": "anonymous/anchor_selection_frequency.csv",
    "Axis_correlations": "anonymous/axis_correlations.csv",
    "Tie_groups": "anonymous/semantic_tie_groups.csv",
}


@dataclass(frozen=True)
class PublicRun:
    repo_root: Path
    run_dir: Path
    private_dir: Path
    config: dict[str, Any]
    manifest: dict[str, Any]
    result_lock: dict[str, Any]
    stage1_rows: list[dict[str, Any]]
    semantic_rows: list[dict[str, Any]]
    anchor_rows: list[dict[str, Any]]
    top50_rows: list[dict[str, Any]]
    anchors: list[dict[str, str]]


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def load_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as error:
        raise ValueError(f"Invalid JSON file: {path}") from error
    if not isinstance(value, dict):
        raise ValueError(f"Expected a JSON object: {path}")
    return value


def load_config(path: Path = CONFIG_PATH) -> dict[str, Any]:
    config = load_json(path)
    if config.get("schema_version") != 1:
        raise ValueError("Unsupported result-data configuration schema.")
    return config


def read_csv(path: Path) -> list[dict[str, str]]:
    try:
        with path.open(newline="", encoding="utf-8") as handle:
            return list(csv.DictReader(handle))
    except (OSError, UnicodeDecodeError, csv.Error) as error:
        raise ValueError(f"Invalid CSV file: {path}") from error


def parse_float(value: Any, label: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{label} must be numeric.") from error
    if not math.isfinite(result):
        raise ValueError(f"{label} must be finite.")
    return result


def parse_int(value: Any, label: str) -> int:
    number = parse_float(value, label)
    if not number.is_integer():
        raise ValueError(f"{label} must be an integer.")
    return int(number)


def parse_bool(value: Any, label: str) -> bool:
    normalized = str(value).strip().lower()
    if normalized not in {"true", "false"}:
        raise ValueError(f"{label} must be True or False.")
    return normalized == "true"


def average_descending_ranks(values: Sequence[float]) -> list[float]:
    order = sorted(range(len(values)), key=lambda index: (-values[index], index))
    ranks = [0.0] * len(values)
    position = 0
    while position < len(order):
        end = position + 1
        value = values[order[position]]
        while end < len(order) and values[order[end]] == value:
            end += 1
        average = ((position + 1) + end) / 2.0
        for offset in range(position, end):
            ranks[order[offset]] = average
        position = end
    return ranks


def _require_regular(path: Path, label: str) -> None:
    if not path.exists() or not path.is_file() or path.is_symlink():
        raise ValueError(f"Required regular file is missing for {label}: {path}")


def _verify_hash(path: Path, expected: str, label: str) -> None:
    _require_regular(path, label)
    actual = sha256_file(path)
    if actual != expected:
        raise ValueError(f"SHA256 mismatch for {label}: {actual}")


def _verify_record(record: Mapping[str, Any], path: Path, label: str) -> None:
    _require_regular(path, label)
    try:
        recorded_path = Path(str(record["path"])).resolve()
        recorded_size = int(record["size"])
        recorded_hash = str(record["sha256"])
    except (KeyError, TypeError, ValueError) as error:
        raise ValueError(f"Malformed locked record for {label}.") from error
    if recorded_path != path.resolve():
        raise ValueError(f"Locked path mismatch for {label}.")
    if path.stat().st_size != recorded_size:
        raise ValueError(f"Locked size mismatch for {label}.")
    _verify_hash(path, recorded_hash, label)


def _axes(config: Mapping[str, Any]) -> list[tuple[str, str]]:
    axes = []
    for item in config["axes"]:
        axes.append((str(item["key"]), str(item["label"])))
    if [key for key, _ in axes] != [
        "ncm811",
        "lithium_metal",
        "lhce_solvation",
    ]:
        raise ValueError("Unexpected fixed axis order.")
    return axes


def _typed_semantic_rows(
    rows: Sequence[Mapping[str, str]],
    axes: Sequence[tuple[str, str]],
) -> list[dict[str, Any]]:
    required = {
        "candidate_id",
        "connectivity_smiles",
        "semantic_score",
        "semantic_rank",
    }
    for axis, _ in axes:
        required.update(
            {
                f"{axis}_score",
                f"{axis}_rank",
                f"{axis}_percentile",
            }
        )
    typed = []
    for index, source in enumerate(rows, start=1):
        if not required.issubset(source):
            raise ValueError("Semantic score columns are incomplete.")
        row: dict[str, Any] = {
            "candidate_id": str(source["candidate_id"]),
            "connectivity_smiles": str(source["connectivity_smiles"]),
        }
        for field in required - {"candidate_id", "connectivity_smiles"}:
            row[field] = parse_float(source[field], f"semantic row {index} {field}")
        typed.append(row)
    return typed


def _typed_stage1_rows(rows: Sequence[Mapping[str, str]]) -> list[dict[str, Any]]:
    typed = []
    for index, source in enumerate(rows, start=1):
        row: dict[str, Any] = {}
        for field, value in source.items():
            if field in {"candidate_id", "connectivity_smiles"}:
                row[field] = str(value)
            elif field == "selected":
                row[field] = parse_bool(value, f"Stage 1 row {index} selected")
            elif field == "stage1_rank":
                row[field] = parse_int(value, f"Stage 1 row {index} rank")
            else:
                row[field] = parse_float(value, f"Stage 1 row {index} {field}")
        typed.append(row)
    return typed


def _typed_anchor_rows(
    rows: Sequence[Mapping[str, str]],
    axes: Sequence[tuple[str, str]],
) -> list[dict[str, Any]]:
    typed = []
    for index, source in enumerate(rows, start=1):
        row: dict[str, Any] = {
            "candidate_id": str(source["candidate_id"]),
            "connectivity_smiles": str(source["connectivity_smiles"]),
        }
        for axis, _ in axes:
            for anchor_index in range(1, 11):
                field = f"{axis}_anchor_A{anchor_index:02d}_cosine"
                row[field] = parse_float(
                    source.get(field), f"anchor row {index} {field}"
                )
        typed.append(row)
    return typed


def _unique_nonempty(rows: Sequence[Mapping[str, Any]], field: str, label: str) -> set[str]:
    values = [str(row.get(field, "")) for row in rows]
    if any(not value for value in values) or len(set(values)) != len(values):
        raise ValueError(f"{label} contains missing or duplicate {field}.")
    return set(values)


def validate_public_run(repo_root: Path, config: Mapping[str, Any]) -> PublicRun:
    repo_root = repo_root.resolve()
    run_id = str(config["run_id"])
    run_dir = repo_root / "runs_v4" / run_id
    private_dir = repo_root / "private_runs_v4" / run_id
    manifest_path = run_dir / "manifest.json"
    lock_path = run_dir / "result.lock.json"
    _verify_hash(
        manifest_path,
        str(config["expected_hashes"]["manifest.json"]),
        "formal manifest",
    )
    _verify_hash(
        lock_path,
        str(config["expected_hashes"]["result.lock.json"]),
        "formal result lock",
    )
    manifest = load_json(manifest_path)
    result_lock = load_json(lock_path)
    if manifest.get("method_version") != config["method_version"]:
        raise ValueError("Formal manifest method version mismatch.")
    if result_lock.get("method_version") != config["method_version"]:
        raise ValueError("Formal result-lock method version mismatch.")
    if result_lock.get("manifest_sha256") != sha256_file(manifest_path):
        raise ValueError("Result lock does not bind the formal manifest.")

    outputs = result_lock.get("outputs")
    if not isinstance(outputs, dict) or set(outputs) != PUBLIC_OUTPUT_RECORDS:
        raise ValueError("Formal result-lock output set is invalid.")
    output_paths = {
        "anchor_cosines.csv": run_dir / "anchor_cosines.csv",
        "diagnostics.json": run_dir / "diagnostics.json",
        "semantic_scores.csv": run_dir / "semantic_scores.csv",
        "semantic_scores.jsonl": run_dir / "semantic_scores.jsonl",
        "top50.csv": run_dir / "top50.csv",
        "semantic_embeddings.npy": run_dir / "semantic_embeddings.npy",
        "embedding_completion.npy": run_dir / "embedding_completion.npy",
        "checkpoint_metadata.json": run_dir / "checkpoints/embedding.metadata.json",
    }
    for key, path in output_paths.items():
        _verify_record(outputs[key], path, f"locked output {key}")

    inputs = manifest.get("inputs")
    if not isinstance(inputs, dict):
        raise ValueError("Formal manifest input set is invalid.")
    input_paths = {
        "stage1_scores": run_dir / "stage1_scores.csv",
        "reference_anchors": run_dir / "blind/reference_anchors.csv",
        "molecular_cards": run_dir / "blind/molecular_cards.csv",
    }
    for key, path in input_paths.items():
        if key not in inputs or not isinstance(inputs[key], dict):
            raise ValueError(f"Formal manifest is missing {key}.")
        _verify_record(inputs[key], path, f"locked input {key}")

    axes = _axes(config)
    counts = config["counts"]
    stage1 = _typed_stage1_rows(read_csv(input_paths["stage1_scores"]))
    semantic = _typed_semantic_rows(
        read_csv(output_paths["semantic_scores.csv"]), axes
    )
    anchor_rows = _typed_anchor_rows(
        read_csv(output_paths["anchor_cosines.csv"]), axes
    )
    top50_raw = read_csv(output_paths["top50.csv"])
    anchors = read_csv(input_paths["reference_anchors"])
    molecular_cards = read_csv(input_paths["molecular_cards"])

    if len(stage1) != int(counts["stage1"]):
        raise ValueError("Stage 1 row count mismatch.")
    if len(semantic) != int(counts["semantic"]):
        raise ValueError("Semantic row count mismatch.")
    if len(anchor_rows) != int(counts["semantic"]):
        raise ValueError("Anchor-cosine row count mismatch.")
    if len(top50_raw) != int(counts["top50"]):
        raise ValueError("Top 50 row count mismatch.")
    if len(anchors) != int(counts["anchors"]):
        raise ValueError("Reference-anchor row count mismatch.")
    if len(molecular_cards) != int(counts["stage1"]):
        raise ValueError("Molecular-card row count mismatch.")

    stage1_ids = _unique_nonempty(stage1, "candidate_id", "Stage 1")
    _unique_nonempty(stage1, "connectivity_smiles", "Stage 1")
    semantic_ids = _unique_nonempty(semantic, "candidate_id", "Semantic scores")
    _unique_nonempty(semantic, "connectivity_smiles", "Semantic scores")
    anchor_ids = _unique_nonempty(anchor_rows, "candidate_id", "Anchor cosines")
    selected_ids = {
        str(row["candidate_id"]) for row in stage1 if bool(row["selected"])
    }
    if len(selected_ids) != int(counts["semantic"]):
        raise ValueError("Stage 1 does not retain exactly 200 candidates.")
    if semantic_ids != selected_ids or anchor_ids != semantic_ids:
        raise ValueError("Stage 1, semantic and anchor candidate sets differ.")
    if not semantic_ids.issubset(stage1_ids):
        raise ValueError("Semantic candidates are absent from Stage 1.")

    expected_anchor_ids = [f"A{index:02d}" for index in range(1, 11)]
    if [row.get("anchor_id") for row in anchors] != expected_anchor_ids:
        raise ValueError("Reference anchors are not ordered A01-A10.")

    tolerance = float(config["float_tolerance"])
    for row in semantic:
        percentiles = []
        for axis, _ in axes:
            rank = float(row[f"{axis}_rank"])
            expected_percentile = (200.0 - rank) / 199.0
            if not math.isclose(
                float(row[f"{axis}_percentile"]),
                expected_percentile,
                rel_tol=0.0,
                abs_tol=tolerance,
            ):
                raise ValueError(f"Axis percentile mismatch for {row['candidate_id']}.")
            percentiles.append(expected_percentile)
        expected_semantic = sum(percentiles) / 3.0
        if not math.isclose(
            float(row["semantic_score"]),
            expected_semantic,
            rel_tol=0.0,
            abs_tol=tolerance,
        ):
            raise ValueError(f"Semantic score mismatch for {row['candidate_id']}.")

    recomputed_semantic_ranks = average_descending_ranks(
        [float(row["semantic_score"]) for row in semantic]
    )
    for row, rank in zip(semantic, recomputed_semantic_ranks):
        if not math.isclose(
            float(row["semantic_rank"]), rank, rel_tol=0.0, abs_tol=tolerance
        ):
            raise ValueError(f"Semantic rank mismatch for {row['candidate_id']}.")

    expected_order = sorted(
        semantic,
        key=lambda row: (-float(row["semantic_score"]), str(row["candidate_id"])),
    )

    semantic_by_id = {str(row["candidate_id"]): row for row in semantic}
    top50_ids = [str(row["candidate_id"]) for row in top50_raw]
    expected_top50_ids = [str(row["candidate_id"]) for row in expected_order[:50]]
    if set(top50_ids) != set(expected_top50_ids):
        raise ValueError("Top 50 candidate set differs from the pure-semantic ranking.")
    for source in top50_raw:
        candidate_id = str(source["candidate_id"])
        reference = semantic_by_id[candidate_id]
        for field in (
            "ncm811_percentile",
            "lithium_metal_percentile",
            "lhce_solvation_percentile",
            "semantic_score",
            "semantic_rank",
        ):
            if not math.isclose(
                parse_float(source[field], f"Top 50 {candidate_id} {field}"),
                float(reference[field]),
                rel_tol=0.0,
                abs_tol=tolerance,
            ):
                raise ValueError(f"Top 50 value mismatch for {candidate_id} {field}.")

    anchor_by_id = {str(row["candidate_id"]): row for row in anchor_rows}
    for candidate_id, semantic_row in semantic_by_id.items():
        source = anchor_by_id[candidate_id]
        for axis, _ in axes:
            values = [
                (
                    f"A{index:02d}",
                    float(source[f"{axis}_anchor_A{index:02d}_cosine"]),
                )
                for index in range(1, 11)
            ]
            top_two = sorted(values, key=lambda item: (-item[1], item[0]))[:2]
            expected_score = sum(value for _, value in top_two) / 2.0
            if not math.isclose(
                expected_score,
                float(semantic_row[f"{axis}_score"]),
                rel_tol=0.0,
                abs_tol=tolerance,
            ):
                raise ValueError(
                    f"Top-2 anchor mean mismatch for {candidate_id} {axis}."
                )

    return PublicRun(
        repo_root=repo_root,
        run_dir=run_dir,
        private_dir=private_dir,
        config=dict(config),
        manifest=manifest,
        result_lock=result_lock,
        stage1_rows=stage1,
        semantic_rows=semantic,
        anchor_rows=anchor_rows,
        top50_rows=top50_raw,
        anchors=[{key: str(value) for key, value in row.items()} for row in anchors],
    )


def _tie_sizes(rows: Sequence[Mapping[str, Any]]) -> dict[float, int]:
    counts: dict[float, int] = {}
    for row in rows:
        score = float(row["semantic_score"])
        counts[score] = counts.get(score, 0) + 1
    return counts


def derive_anonymous_tables(public: PublicRun) -> dict[str, Any]:
    axes = _axes(public.config)
    anchor_names = {
        str(row["anchor_id"]): str(row["name"]) for row in public.anchors
    }
    tie_sizes = _tie_sizes(public.semantic_rows)
    ranking_all: list[dict[str, Any]] = []
    ordered_sources = sorted(
        public.semantic_rows,
        key=lambda row: (-float(row["semantic_score"]), str(row["candidate_id"])),
    )
    for final_rank, source in enumerate(ordered_sources, start=1):
        score = float(source["semantic_score"])
        row = {
            "candidate_id": source["candidate_id"],
            "connectivity_smiles": source["connectivity_smiles"],
        }
        for axis, _ in axes:
            for suffix in ("score", "rank", "percentile"):
                row[f"{axis}_{suffix}"] = float(source[f"{axis}_{suffix}"])
        row.update(
            {
                "semantic_score": score,
                "semantic_rank": float(source["semantic_rank"]),
                "final_rank": final_rank,
                "is_top20": final_rank <= 20,
                "is_top50": final_rank <= 50,
                "semantic_tie_size": tie_sizes[score],
                "candidate_id_tiebreak_applied": tie_sizes[score] > 1,
            }
        )
        ranking_all.append(row)

    axis_scores_long: list[dict[str, Any]] = []
    for row in ranking_all:
        for axis, label in axes:
            axis_scores_long.append(
                {
                    "candidate_id": row["candidate_id"],
                    "connectivity_smiles": row["connectivity_smiles"],
                    "axis": axis,
                    "axis_label": label,
                    "axis_score": row[f"{axis}_score"],
                    "axis_rank": row[f"{axis}_rank"],
                    "axis_percentile": row[f"{axis}_percentile"],
                    "semantic_score": row["semantic_score"],
                    "semantic_rank": row["semantic_rank"],
                    "final_rank": row["final_rank"],
                }
            )

    raw_anchor_by_id = {
        str(row["candidate_id"]): row for row in public.anchor_rows
    }
    anchor_cosines_long: list[dict[str, Any]] = []
    for row in ranking_all:
        source = raw_anchor_by_id[str(row["candidate_id"])]
        for axis, label in axes:
            values = [
                (
                    f"A{index:02d}",
                    float(source[f"{axis}_anchor_A{index:02d}_cosine"]),
                )
                for index in range(1, 11)
            ]
            ordered = sorted(values, key=lambda item: (-item[1], item[0]))
            top_order = {anchor_id: position for position, (anchor_id, _) in enumerate(ordered[:2], start=1)}
            for anchor_id, value in values:
                anchor_cosines_long.append(
                    {
                        "candidate_id": row["candidate_id"],
                        "connectivity_smiles": row["connectivity_smiles"],
                        "final_rank": row["final_rank"],
                        "axis": axis,
                        "axis_label": label,
                        "anchor_id": anchor_id,
                        "anchor_name": anchor_names[anchor_id],
                        "cosine_similarity": value,
                        "top2_order": top_order.get(anchor_id, ""),
                        "is_top2": anchor_id in top_order,
                    }
                )

    top20_axis_matrix = []
    top20_fields = [
        "candidate_id",
        "connectivity_smiles",
        "final_rank",
        "ncm811_score",
        "ncm811_rank",
        "ncm811_percentile",
        "lithium_metal_score",
        "lithium_metal_rank",
        "lithium_metal_percentile",
        "lhce_solvation_score",
        "lhce_solvation_rank",
        "lhce_solvation_percentile",
        "semantic_score",
        "semantic_rank",
    ]
    for row in ranking_all[:20]:
        top20_axis_matrix.append({field: row[field] for field in top20_fields})
    top50_ranking = [dict(row) for row in ranking_all[:50]]

    stage1_all = sorted(public.stage1_rows, key=lambda row: int(row["stage1_rank"]))

    correlations = []
    axis_values = {
        axis: np.asarray(
            [float(row[f"{axis}_score"]) for row in ranking_all],
            dtype=np.float64,
        )
        for axis, _ in axes
    }
    axis_ranks = {
        axis: np.asarray(average_descending_ranks(values), dtype=np.float64)
        for axis, values in axis_values.items()
    }
    for method in ("pearson", "spearman"):
        for left, _ in axes:
            for right, _ in axes:
                left_values = axis_values[left] if method == "pearson" else axis_ranks[left]
                right_values = (
                    axis_values[right] if method == "pearson" else axis_ranks[right]
                )
                correlations.append(
                    {
                        "method": method,
                        "axis_1": left,
                        "axis_2": right,
                        "correlation": float(np.corrcoef(left_values, right_values)[0, 1]),
                    }
                )

    axis_distribution = []
    for axis, label in axes:
        values = np.asarray(
            [float(row[f"{axis}_score"]) for row in ranking_all],
            dtype=np.float64,
        )
        quantiles = np.quantile(values, [0.0, 0.05, 0.25, 0.5, 0.75, 0.95, 1.0])
        axis_distribution.append(
            {
                "axis": axis,
                "axis_label": label,
                "count": int(values.size),
                "minimum": float(quantiles[0]),
                "p05": float(quantiles[1]),
                "p25": float(quantiles[2]),
                "median": float(quantiles[3]),
                "mean": float(values.mean()),
                "standard_deviation": float(values.std(ddof=0)),
                "p75": float(quantiles[4]),
                "p95": float(quantiles[5]),
                "maximum": float(quantiles[6]),
            }
        )

    frequency_counts: dict[tuple[str, str], int] = {}
    for row in anchor_cosines_long:
        if row["is_top2"]:
            key = (str(row["axis"]), str(row["anchor_id"]))
            frequency_counts[key] = frequency_counts.get(key, 0) + 1
    anchor_frequency = []
    for axis, label in axes:
        for anchor_id in [f"A{index:02d}" for index in range(1, 11)]:
            count = frequency_counts.get((axis, anchor_id), 0)
            anchor_frequency.append(
                {
                    "axis": axis,
                    "axis_label": label,
                    "anchor_id": anchor_id,
                    "anchor_name": anchor_names[anchor_id],
                    "selection_count": count,
                    "selection_fraction_of_slots": count / 400.0,
                    "candidate_fraction": count / 200.0,
                }
            )

    tie_groups = []
    tied_scores = sorted(
        (score for score, size in tie_sizes.items() if size > 1), reverse=True
    )
    for group_index, score in enumerate(tied_scores, start=1):
        members = sorted(
            (row for row in ranking_all if float(row["semantic_score"]) == score),
            key=lambda row: int(row["final_rank"]),
        )
        for order, row in enumerate(members, start=1):
            tie_groups.append(
                {
                    "tie_group_id": f"T{group_index:02d}",
                    "semantic_score": score,
                    "semantic_rank": row["semantic_rank"],
                    "tie_size": len(members),
                    "order_within_tie": order,
                    "candidate_id": row["candidate_id"],
                    "final_rank": row["final_rank"],
                    "deterministic_tie_rule": "candidate_id ascending",
                }
            )

    processing_summary = {
        "run_id": public.config["run_id"],
        "method_version": public.config["method_version"],
        "candidate_count_stage1": len(stage1_all),
        "candidate_count_semantic": len(ranking_all),
        "top50_count": len(top50_ranking),
        "axis_score_count": len(axis_scores_long),
        "anchor_cosine_count": len(anchor_cosines_long),
        "semantic_tie_group_count": len(tied_scores),
        "candidates_in_semantic_ties": len(tie_groups),
        "maximum_semantic_tie_size": max(
            (size for size in tie_sizes.values()), default=1
        ),
        "top_candidate_id": ranking_all[0]["candidate_id"],
        "top_candidate_semantic_tie_size": ranking_all[0]["semantic_tie_size"],
        "top_candidate_candidate_id_tiebreak_applied": ranking_all[0][
            "candidate_id_tiebreak_applied"
        ],
        "identity_used_for_ranking": False,
    }
    if processing_summary["semantic_tie_group_count"] != 24:
        raise ValueError("Unexpected formal semantic tie-group count.")
    if processing_summary["candidates_in_semantic_ties"] != 52:
        raise ValueError("Unexpected formal tied-candidate count.")
    if processing_summary["maximum_semantic_tie_size"] != 3:
        raise ValueError("Unexpected formal maximum semantic tie size.")

    return {
        "ranking_all": ranking_all,
        "axis_scores_long": axis_scores_long,
        "anchor_cosines_long": anchor_cosines_long,
        "top20_axis_matrix": top20_axis_matrix,
        "top50_ranking": top50_ranking,
        "stage1_all": stage1_all,
        "axis_correlations": correlations,
        "axis_distribution_summary": axis_distribution,
        "anchor_selection_frequency": anchor_frequency,
        "semantic_tie_groups": tie_groups,
        "processing_summary": processing_summary,
    }


def annotate_selected_candidate(
    public: PublicRun,
    anonymous: Mapping[str, Any],
) -> dict[str, Any]:
    ranking = anonymous["ranking_all"]
    if not ranking or int(ranking[0]["final_rank"]) != 1:
        raise ValueError("Anonymous ranking has no deterministic first candidate.")
    selected = ranking[0]
    expected = public.config["selected_identity"]
    if selected["candidate_id"] != expected["expected_candidate_id"]:
        raise ValueError("Public rank-1 candidate differs from the fixed formal result.")
    if selected["connectivity_smiles"] != expected["expected_connectivity_smiles"]:
        raise ValueError("Public rank-1 structure differs from the fixed formal result.")

    registry_path = public.private_dir / "target_registry.json"
    manifest_hashes = public.manifest.get("private_input_hashes")
    if not isinstance(manifest_hashes, dict) or "target_registry" not in manifest_hashes:
        raise ValueError("Formal manifest lacks the private identity hash.")
    _verify_hash(
        registry_path,
        str(manifest_hashes["target_registry"]),
        "private identity registry",
    )
    registry = load_json(registry_path)
    if registry.get("candidate_id") != selected["candidate_id"]:
        raise ValueError("Identity registry does not describe the public rank-1 candidate.")
    if registry.get("connectivity_smiles") != selected["connectivity_smiles"]:
        raise ValueError("Identity registry structure differs from public rank 1.")

    reveal_path = public.private_dir / "target_reveal.json"
    _require_regular(reveal_path, "formal identity annotation")
    reveal = load_json(reveal_path)
    if (
        reveal.get("candidate_id") != selected["candidate_id"]
        or reveal.get("manifest_sha256")
        != sha256_file(public.run_dir / "manifest.json")
        or reveal.get("result_lock_sha256")
        != sha256_file(public.run_dir / "result.lock.json")
    ):
        raise ValueError("Formal identity annotation is not bound to the locked result.")

    stage1 = next(
        row
        for row in public.stage1_rows
        if row["candidate_id"] == selected["candidate_id"]
    )
    summary: dict[str, Any] = {
        "candidate_id": selected["candidate_id"],
        "display_name": expected["display_name"],
        "systematic_name": expected["systematic_name"],
        "connectivity_smiles": selected["connectivity_smiles"],
        "source_canonical_smiles": registry.get("source_canonical_smiles"),
        "molecular_formula": registry.get("molecular_formula"),
        "molecular_weight": parse_float(
            registry.get("molecular_weight"), "selected molecular weight"
        ),
        "cid": parse_int(registry.get("cid"), "selected PubChem CID"),
        "cas": str(registry.get("cas")),
        "inchi_key": str(registry.get("inchi_key")),
        "stage1_rank": int(stage1["stage1_rank"]),
    }
    for axis, _ in _axes(public.config):
        summary[f"{axis}_score"] = selected[f"{axis}_score"]
        summary[f"{axis}_rank"] = selected[f"{axis}_rank"]
        summary[f"{axis}_percentile"] = selected[f"{axis}_percentile"]
    summary.update(
        {
            "semantic_score": selected["semantic_score"],
            "semantic_rank": selected["semantic_rank"],
            "final_rank": selected["final_rank"],
            "semantic_tie_size": selected["semantic_tie_size"],
            "candidate_id_tiebreak_applied": selected[
                "candidate_id_tiebreak_applied"
            ],
            "identity_annotation_stage": "after_anonymous_final_ranking",
            "identity_used_for_ranking": False,
        }
    )

    annotated_ranking = []
    for row in ranking:
        is_selected = row["candidate_id"] == selected["candidate_id"]
        annotated_ranking.append(
            {
                **row,
                "is_selected_candidate": is_selected,
                "display_name": expected["display_name"] if is_selected else "",
                "systematic_name": (
                    expected["systematic_name"] if is_selected else ""
                ),
            }
        )
    selected_anchor = []
    for row in anonymous["anchor_cosines_long"]:
        if row["candidate_id"] == selected["candidate_id"]:
            selected_anchor.append(
                {
                    **row,
                    "display_name": expected["display_name"],
                }
            )
    if len(selected_anchor) != 30:
        raise ValueError("Selected candidate anchor profile must have 30 rows.")
    for axis, _ in _axes(public.config):
        top_two = [
            row["anchor_id"]
            for row in selected_anchor
            if row["axis"] == axis and row["is_top2"]
        ]
        if top_two != ["A05", "A06"]:
            raise ValueError(f"Unexpected selected-candidate Top-2 anchors for {axis}.")

    return {
        "selected_candidate_summary": summary,
        "ranking_all_annotated": annotated_ranking,
        "selected_anchor_profile": selected_anchor,
    }


def _format_csv_value(value: Any) -> Any:
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("Cannot serialize a non-finite float.")
        return format(value, ".17g")
    return value


def write_csv(path: Path, rows: Sequence[Mapping[str, Any]]) -> None:
    if not rows:
        raise ValueError(f"Cannot write an empty CSV: {path.name}")
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0])
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            if list(row) != fields:
                raise ValueError(f"Inconsistent CSV fields for {path.name}.")
            writer.writerow({key: _format_csv_value(row[key]) for key in fields})


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    text = (
        json.dumps(
            value,
            indent=2,
            sort_keys=True,
            ensure_ascii=False,
            allow_nan=False,
        )
        + "\n"
    )
    path.write_text(text, encoding="utf-8")


def _typed_csv_cell(value: str) -> Any:
    if value == "":
        return None
    if value in {"True", "False"}:
        return value == "True"
    if re.fullmatch(r"[+-]?\d+", value):
        return int(value)
    if re.fullmatch(
        r"[+-]?(?:\d+\.\d*|\d*\.\d+|\d+)(?:[eE][+-]?\d+)?", value
    ):
        return float(value)
    return value


def _read_typed_csv(path: Path) -> tuple[list[str], list[list[Any]]]:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        rows = [[_typed_csv_cell(value) for value in row] for row in reader]
    return header, rows


def _require_openpyxl():
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError(
            "openpyxl==3.1.5 is required; install the figure requirements."
        ) from error
    if openpyxl.__version__ != "3.1.5":
        raise RuntimeError(
            f"Expected openpyxl 3.1.5, found {openpyxl.__version__}."
        )
    return openpyxl


def build_workbook(data_dir: Path, public: PublicRun) -> Path:
    openpyxl = _require_openpyxl()
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "Organic Additive Semantic Fingerprint"
    workbook.properties.title = "Semantic screening source data"
    created = datetime.fromisoformat(str(public.result_lock["created_at"]))
    workbook.properties.created = created
    workbook.properties.modified = created

    readme = workbook.create_sheet("README")
    readme.sheet_view.showGridLines = False
    readme_rows = [
        ["Semantic screening source data", None],
        ["Run ID", public.config["run_id"]],
        ["Method version", public.config["method_version"]],
        ["Purpose", "Source data for the semantic screening Results section"],
        [
            "Ranking order",
            "semantic_score descending; exact ties by anonymous candidate_id ascending",
        ],
        [
            "Identity boundary",
            "The selected identity is annotated only after anonymous final_rank = 1 is fixed.",
        ],
        ["Stage 1 rows", 800],
        ["Semantic candidates", 200],
        ["Reported Top N", 50],
        ["Canonical data", "CSV and JSON files beside this workbook"],
        [
            "Numeric precision",
            "Full numeric values are stored; worksheet display is limited to six decimals.",
        ],
    ]
    for row in readme_rows:
        readme.append(row)
    readme["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    readme["A1"].fill = PatternFill("solid", fgColor="27313A")
    readme["B1"].fill = PatternFill("solid", fgColor="27313A")
    readme["A1"].alignment = Alignment(vertical="center")
    readme.row_dimensions[1].height = 26
    for row in range(2, readme.max_row + 1):
        readme.cell(row, 1).font = Font(name="Arial", size=10, bold=True)
        readme.cell(row, 2).font = Font(name="Arial", size=10)
        readme.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 90
    readme.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="27313A")
    selected_fill = PatternFill("solid", fgColor="E8F1F5")
    header_font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=9)
    thin_bottom = Border(bottom=Side(style="thin", color="D9DEE3"))

    for sheet_name, relative_path in WORKBOOK_SHEETS.items():
        header, rows = _read_typed_csv(data_dir / relative_path)
        sheet = workbook.create_sheet(sheet_name)
        sheet.sheet_view.showGridLines = False
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        sheet.row_dimensions[1].height = 30
        for row_index in range(2, sheet.max_row + 1):
            selected = False
            if "is_selected_candidate" in header:
                flag_col = header.index("is_selected_candidate") + 1
                selected = sheet.cell(row_index, flag_col).value is True
            for cell in sheet[row_index]:
                cell.font = Font(
                    name="Arial",
                    size=9,
                    bold=selected,
                )
                cell.border = thin_bottom
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
                elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                    cell.number_format = "0"
                if selected:
                    cell.fill = selected_fill
        for column_index, field in enumerate(header, start=1):
            values = [str(field)]
            for row_index in range(2, min(sheet.max_row, 80) + 1):
                value = sheet.cell(row_index, column_index).value
                values.append("" if value is None else str(value))
            width = max(len(value) for value in values) + 2
            if "smiles" in field:
                width = min(max(width, 24), 44)
            else:
                width = min(max(width, 10), 28)
            sheet.column_dimensions[
                openpyxl.utils.get_column_letter(column_index)
            ].width = width

    output = data_dir / "semantic_screening_source_data.xlsx"
    workbook.save(output)
    return output


def _workbook_value_equal(left: Any, right: Any, tolerance: float) -> bool:
    if left is None and right is None:
        return True
    if isinstance(left, bool) or isinstance(right, bool):
        return left is right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(
            float(left), float(right), rel_tol=0.0, abs_tol=tolerance
        )
    return str(left) == str(right)


def verify_workbook(data_dir: Path, tolerance: float) -> None:
    openpyxl = _require_openpyxl()
    workbook_path = data_dir / "semantic_screening_source_data.xlsx"
    workbook = openpyxl.load_workbook(workbook_path, data_only=False, read_only=False)
    expected_sheets = ["README", *WORKBOOK_SHEETS]
    if workbook.sheetnames != expected_sheets:
        raise ValueError("Source Data workbook sheet set or order is invalid.")
    for sheet_name, relative_path in WORKBOOK_SHEETS.items():
        header, rows = _read_typed_csv(data_dir / relative_path)
        sheet = workbook[sheet_name]
        if sheet.max_row != len(rows) + 1 or sheet.max_column != len(header):
            raise ValueError(f"Workbook dimensions differ for {sheet_name}.")
        for column, expected in enumerate(header, start=1):
            if sheet.cell(1, column).value != expected:
                raise ValueError(f"Workbook header differs for {sheet_name}.")
        for row_index, source_row in enumerate(rows, start=2):
            for column_index, expected in enumerate(source_row, start=1):
                actual = sheet.cell(row_index, column_index).value
                if not _workbook_value_equal(actual, expected, tolerance):
                    raise ValueError(
                        f"Workbook value mismatch at {sheet_name} "
                        f"R{row_index}C{column_index}."
                    )
                if sheet.cell(row_index, column_index).data_type == "f":
                    raise ValueError("Source Data workbook must not contain formulas.")
                if isinstance(actual, str) and actual.startswith(
                    ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
                ):
                    raise ValueError("Workbook contains an Excel error value.")
    workbook.close()


def _git_value(repo_root: Path, *args: str) -> str:
    try:
        return subprocess.check_output(
            ["git", *args],
            cwd=repo_root,
            text=True,
            stderr=subprocess.DEVNULL,
        ).strip()
    except (OSError, subprocess.CalledProcessError):
        return "unavailable"


def _record(path: Path, root: Path) -> dict[str, Any]:
    return {
        "path": str(path.relative_to(root)),
        "sha256": sha256_file(path),
        "size": path.stat().st_size,
    }


def _write_generated(
    staging: Path,
    public: PublicRun,
    anonymous: Mapping[str, Any],
    annotated: Mapping[str, Any],
) -> None:
    source_dir = staging / "source_snapshot"
    source_dir.mkdir(parents=True)
    for destination_name, relative_source in public.config["source_snapshot"].items():
        source = public.run_dir / str(relative_source)
        _require_regular(source, f"snapshot source {destination_name}")
        (source_dir / destination_name).write_bytes(source.read_bytes())

    anonymous_dir = staging / "anonymous"
    annotated_dir = staging / "annotated"
    write_csv(anonymous_dir / "ranking_all.csv", anonymous["ranking_all"])
    write_csv(
        anonymous_dir / "axis_scores_long.csv", anonymous["axis_scores_long"]
    )
    write_csv(
        anonymous_dir / "anchor_cosines_long.csv",
        anonymous["anchor_cosines_long"],
    )
    write_csv(
        anonymous_dir / "top20_axis_matrix.csv",
        anonymous["top20_axis_matrix"],
    )
    write_csv(
        anonymous_dir / "top50_ranking.csv", anonymous["top50_ranking"]
    )
    write_csv(anonymous_dir / "stage1_all.csv", anonymous["stage1_all"])
    write_csv(
        anonymous_dir / "axis_correlations.csv",
        anonymous["axis_correlations"],
    )
    write_csv(
        anonymous_dir / "axis_distribution_summary.csv",
        anonymous["axis_distribution_summary"],
    )
    write_csv(
        anonymous_dir / "anchor_selection_frequency.csv",
        anonymous["anchor_selection_frequency"],
    )
    write_csv(
        anonymous_dir / "semantic_tie_groups.csv",
        anonymous["semantic_tie_groups"],
    )
    write_json(
        anonymous_dir / "processing_summary.json",
        anonymous["processing_summary"],
    )

    write_json(
        annotated_dir / "selected_candidate_summary.json",
        annotated["selected_candidate_summary"],
    )
    write_csv(
        annotated_dir / "ranking_all_annotated.csv",
        annotated["ranking_all_annotated"],
    )
    write_csv(
        annotated_dir / "selected_anchor_profile.csv",
        annotated["selected_anchor_profile"],
    )
    build_workbook(staging, public)

    output_records = {}
    for path in sorted(staging.rglob("*")):
        if path.is_file() and path.name != "data_manifest.json":
            output_records[str(path.relative_to(staging))] = _record(path, staging)
    try:
        import openpyxl

        openpyxl_version = openpyxl.__version__
    except ImportError:
        openpyxl_version = "unavailable"
    manifest = {
        "schema_version": 1,
        "run_id": public.config["run_id"],
        "method_version": public.config["method_version"],
        "created_at": utc_now(),
        "formal_manifest_sha256": sha256_file(public.run_dir / "manifest.json"),
        "formal_result_lock_sha256": sha256_file(
            public.run_dir / "result.lock.json"
        ),
        "processing_code": {
            "path": str(Path(__file__).resolve().relative_to(public.repo_root)),
            "sha256": sha256_file(Path(__file__).resolve()),
            "git_commit": _git_value(public.repo_root, "rev-parse", "HEAD"),
        },
        "runtime": {
            "python": sys.version.split()[0],
            "numpy": np.__version__,
            "openpyxl": openpyxl_version,
        },
        "fixed_parameters": {
            "main_top_n": int(public.config["counts"]["main_top_n"]),
            "reported_top_n": int(public.config["counts"]["top50"]),
            "axis_order": [key for key, _ in _axes(public.config)],
            "float_tolerance": float(public.config["float_tolerance"]),
        },
        "selection": {
            "anonymous_rank1_candidate_id": anonymous["ranking_all"][0][
                "candidate_id"
            ],
            "identity_annotation_stage": "after_anonymous_final_ranking",
            "identity_used_for_ranking": False,
        },
        "outputs": output_records,
    }
    write_json(staging / "data_manifest.json", manifest)


def _safe_output_dir(output_dir: Path, repo_root: Path) -> Path:
    output = output_dir.resolve()
    repo = repo_root.resolve()
    forbidden = {
        repo,
        (repo / "runs_v4").resolve(),
        (repo / "private_runs_v4").resolve(),
        Path.home().resolve(),
        Path("/").resolve(),
    }
    if output in forbidden:
        raise ValueError(f"Unsafe generated-data output directory: {output}")
    for protected in ((repo / "runs_v4").resolve(), (repo / "private_runs_v4").resolve()):
        if protected in output.parents:
            raise ValueError("Generated figure data cannot be written into run roots.")
    return output


def verify_generated(
    repo_root: Path,
    output_dir: Path,
    config: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    config = dict(config or load_config())
    public = validate_public_run(repo_root, config)
    output_dir = _safe_output_dir(output_dir, repo_root)
    manifest_path = output_dir / "data_manifest.json"
    manifest = load_json(manifest_path)
    if (
        manifest.get("schema_version") != 1
        or manifest.get("run_id") != config["run_id"]
        or manifest.get("method_version") != config["method_version"]
    ):
        raise ValueError("Generated data manifest boundary is invalid.")
    if manifest.get("formal_manifest_sha256") != sha256_file(
        public.run_dir / "manifest.json"
    ):
        raise ValueError("Generated data is bound to a different formal manifest.")
    if manifest.get("formal_result_lock_sha256") != sha256_file(
        public.run_dir / "result.lock.json"
    ):
        raise ValueError("Generated data is bound to a different result lock.")
    outputs = manifest.get("outputs")
    if not isinstance(outputs, dict):
        raise ValueError("Generated output manifest is malformed.")
    for relative, record in outputs.items():
        path = output_dir / str(relative)
        _require_regular(path, f"generated output {relative}")
        if path.stat().st_size != int(record["size"]):
            raise ValueError(f"Generated output size mismatch: {relative}")
        if sha256_file(path) != str(record["sha256"]):
            raise ValueError(f"Generated output hash mismatch: {relative}")

    expected_counts = {
        "anonymous/ranking_all.csv": 200,
        "anonymous/axis_scores_long.csv": 600,
        "anonymous/anchor_cosines_long.csv": 6000,
        "anonymous/top20_axis_matrix.csv": 20,
        "anonymous/top50_ranking.csv": 50,
        "anonymous/stage1_all.csv": 800,
        "annotated/selected_anchor_profile.csv": 30,
    }
    for relative, expected in expected_counts.items():
        if len(read_csv(output_dir / relative)) != expected:
            raise ValueError(f"Generated row count mismatch: {relative}")
    selected = load_json(output_dir / "annotated/selected_candidate_summary.json")
    expected_identity = config["selected_identity"]
    fixed_values = {
        "candidate_id": expected_identity["expected_candidate_id"],
        "display_name": expected_identity["display_name"],
        "stage1_rank": 141,
        "ncm811_rank": 1.0,
        "lithium_metal_rank": 1.0,
        "lhce_solvation_rank": 1.0,
        "semantic_rank": 1.0,
        "final_rank": 1,
        "candidate_id_tiebreak_applied": False,
    }
    for field, expected in fixed_values.items():
        if selected.get(field) != expected:
            raise ValueError(f"Selected-candidate acceptance mismatch: {field}")
    verify_workbook(output_dir, float(config["float_tolerance"]))
    return manifest


def prepare(
    repo_root: Path,
    output_dir: Path,
    replace_generated: bool = False,
) -> dict[str, Any]:
    config = load_config()
    repo_root = repo_root.resolve()
    output_dir = _safe_output_dir(output_dir, repo_root)
    if output_dir.exists() and not replace_generated:
        try:
            manifest = verify_generated(repo_root, output_dir, config)
        except Exception as error:
            raise FileExistsError(
                "Generated data already exists but does not verify; "
                "use --replace-generated only after reviewing the mismatch."
            ) from error
        return {"status": "unchanged", "manifest": manifest}
    if output_dir.exists() and replace_generated:
        existing_manifest = output_dir / "data_manifest.json"
        if not existing_manifest.exists():
            raise ValueError(
                "Refusing to replace a directory without a generated data manifest."
            )
        existing = load_json(existing_manifest)
        if (
            existing.get("schema_version") != 1
            or existing.get("run_id") != config["run_id"]
        ):
            raise ValueError("Refusing to replace an unrelated generated directory.")

    public = validate_public_run(repo_root, config)
    anonymous = derive_anonymous_tables(public)
    # This call is intentionally after every anonymous table has been derived.
    annotated = annotate_selected_candidate(public, anonymous)

    output_dir.parent.mkdir(parents=True, exist_ok=True)
    staging = output_dir.parent / f".{output_dir.name}.tmp-{uuid.uuid4().hex}"
    staging.mkdir()
    backup: Path | None = None
    try:
        _write_generated(staging, public, anonymous, annotated)
        verify_generated(repo_root, staging, config)
        if output_dir.exists():
            backup = output_dir.parent / f".{output_dir.name}.backup-{uuid.uuid4().hex}"
            output_dir.replace(backup)
        staging.replace(output_dir)
        if backup is not None:
            shutil.rmtree(backup)
    except Exception:
        if staging.exists():
            shutil.rmtree(staging)
        if backup is not None and backup.exists() and not output_dir.exists():
            backup.replace(output_dir)
        raise
    manifest = verify_generated(repo_root, output_dir, config)
    return {"status": "created", "manifest": manifest}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Prepare and verify semantic-screening publication data."
    )
    subparsers = parser.add_subparsers(dest="command", required=True)
    for command in ("prepare", "verify"):
        subparser = subparsers.add_parser(command)
        subparser.add_argument(
            "--repo-root",
            type=Path,
            default=DEFAULT_REPO_ROOT,
        )
        subparser.add_argument(
            "--output-dir",
            type=Path,
            default=DEFAULT_OUTPUT_DIR,
        )
        if command == "prepare":
            subparser.add_argument(
                "--replace-generated",
                action="store_true",
            )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        if args.command == "prepare":
            result = prepare(
                args.repo_root,
                args.output_dir,
                args.replace_generated,
            )
        else:
            result = {
                "status": "verified",
                "manifest": verify_generated(args.repo_root, args.output_dir),
            }
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1
    manifest = result["manifest"]
    print(
        json.dumps(
            {
                "status": result["status"],
                "run_id": manifest["run_id"],
                "output_dir": str(args.output_dir.resolve()),
                "selected_candidate_id": manifest["selection"][
                    "anonymous_rank1_candidate_id"
                ],
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
